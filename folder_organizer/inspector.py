"""File Inspector — MIME detection, text extraction, SHA-256 hashing."""

from __future__ import annotations

import hashlib
import mimetypes
from pathlib import Path
from typing import Optional

import docx
import openpyxl
import pypdf
from pptx import Presentation as pptx

from folder_organizer.models import ExtractionStatus, InspectedFile

# Attempt to import python-magic; set to None if unavailable so callers can
# detect the absence and fall back to mimetypes.
try:
    import magic
except ImportError:
    magic = None  # type: ignore[assignment]

_MAX_CHARS = 4096


class FileInspector:
    """Inspect a single file: compute its SHA-256 hash, detect its MIME type,
    and extract a text snippet where possible."""

    def inspect(self, path: Path) -> InspectedFile:
        """Return an :class:`InspectedFile` for *path*.

        SHA-256 is always computed when the file is readable.  Text extraction
        is attempted based on MIME type and file extension; failures are
        recorded in ``extraction_status`` without re-raising.  If even the
        initial file read fails (e.g. permission/I/O error), the method returns
        an :class:`InspectedFile` with ``extraction_status = FAILED`` and
        placeholder SHA-256 / size values rather than propagating the error.
        """
        try:
            # -------------------------------------------------------------- #
            # 1. SHA-256 and size (computed before extraction)
            # -------------------------------------------------------------- #
            raw_content = path.read_bytes()
            sha256 = hashlib.sha256(raw_content).hexdigest()
            size_bytes = path.stat().st_size
        except Exception:
            # Catastrophic read failure — return a minimal FAILED record.
            return InspectedFile(
                path=path,
                sha256="",
                size_bytes=0,
                mime_type="application/octet-stream",
                extracted_text=None,
                char_count=None,
                extraction_status=ExtractionStatus.FAILED,
            )

        # ------------------------------------------------------------------ #
        # 2. MIME type detection
        # ------------------------------------------------------------------ #
        mime_type = _detect_mime(path)

        # ------------------------------------------------------------------ #
        # 3. Text extraction dispatch
        # ------------------------------------------------------------------ #
        extracted_text: Optional[str] = None
        char_count: Optional[int] = None
        extraction_status = ExtractionStatus.UNAVAILABLE

        try:
            text, status = _extract(path, mime_type, raw_content)
            extraction_status = status
            if status == ExtractionStatus.OK and text is not None:
                extracted_text = text
                char_count = len(text)
        except Exception:
            extraction_status = ExtractionStatus.FAILED
            extracted_text = None
            char_count = None

        return InspectedFile(
            path=path,
            sha256=sha256,
            size_bytes=size_bytes,
            mime_type=mime_type,
            extracted_text=extracted_text,
            char_count=char_count,
            extraction_status=extraction_status,
        )


# --------------------------------------------------------------------------- #
# Private helpers
# --------------------------------------------------------------------------- #

def _detect_mime(path: Path) -> str:
    """Detect the MIME type of *path*.

    Priority:
    1. ``python-magic`` content-based detection.
    2. ``mimetypes.guess_type`` extension-based fallback.
    3. ``"application/octet-stream"`` if both methods fail.
    """
    if magic is not None:
        try:
            result = magic.from_file(str(path), mime=True)
            if result:
                return result
        except Exception:
            pass

    guessed, _ = mimetypes.guess_type(str(path))
    if guessed:
        return guessed

    return "application/octet-stream"


def _extract(
    path: Path, mime_type: str, raw_content: bytes
) -> tuple[Optional[str], ExtractionStatus]:
    """Dispatch to the appropriate extractor.

    Returns ``(text, status)``.  Raises on error so the caller can catch and
    set ``FAILED``.
    """
    suffix = path.suffix.lower()

    # ---- text/* ----------------------------------------------------------
    if mime_type.startswith("text/"):
        text = raw_content.decode("utf-8", errors="replace")[:_MAX_CHARS]
        return text, ExtractionStatus.OK

    # ---- PDF -------------------------------------------------------------
    if mime_type == "application/pdf":
        reader = pypdf.PdfReader(path)
        pages = reader.pages[:10]
        text = "\n".join(
            page.extract_text() or "" for page in pages
        )[:_MAX_CHARS]
        return text, ExtractionStatus.OK

    # ---- DOCX ------------------------------------------------------------
    if suffix == ".docx":
        doc = docx.Document(path)
        text = " ".join(p.text for p in doc.paragraphs)[:_MAX_CHARS]
        return text, ExtractionStatus.OK

    # ---- XLSX ------------------------------------------------------------
    if suffix == ".xlsx":
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        parts: list[str] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        parts.append(str(cell.value))
        text = " ".join(parts)[:_MAX_CHARS]
        wb.close()
        return text, ExtractionStatus.OK

    # ---- PPTX ------------------------------------------------------------
    if suffix == ".pptx":
        prs = pptx(path)
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            parts.append(run.text)
        text = " ".join(parts)[:_MAX_CHARS]
        return text, ExtractionStatus.OK

    # ---- Everything else -------------------------------------------------
    return None, ExtractionStatus.UNAVAILABLE
